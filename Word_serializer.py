#Az kapott adatott itt konvertáljuk át egy word alapú fájlba.
from docx import Document
from docx.shared import Cm
#Kezdeti teszt adatok később figyelni kell a serializálásra hogy  jól adjuk vissza lehet hogy
# jobb lenne nem kulcs indexek alapján pharsolnám be.
#ERRE MINDENKÉPPEN FIGYLEJ ODA A PANDAS BEKÖTÉSNÉL.

# Az elsőkét oszlop a systemid Gépüzem részére nem szükslégesek.
# Rendszermegnevzés fajták:
test_data = {
        #Egy oszlopos cellák:
        "rendszer_azonositoja_es_megnevezese":"lsd. csatolt",
        "rendszer_tulajdonos":"Pinczés Zoltán",
        "munkahelyi_vezeto": "Gulyás Csaba Lajos",
        "felhasznalo_neve":"Kis Pista",# Változó 1
        "felhasznalo_feladatkore a rendszerben": "Operátor",# lenyíló lista amelyhez lehet adni  vagy törölni, állandó értéken mozog.Ez az excellben lesz bent.

        "torzs_szama":23456, #Főkulcs adat.
        "login_nev":"Felhasznalónév",# Adott.
        "szervezeti_egyseg":"RGK",#AD Szerint.

        "uj_felhasznalo":False,#
        "modositas":False,#
        "visszavonas":False,

        "elektronikus_felelosseg":False,
        "igénylo_azonositoja": "RGKCORP_KIS_XXX",#Adott RGKCORP-KIS-XXX növekményesen fusson. A lépéseket könyveljük le tehát folyamtosan kell növelni.
        "telefon": 36_20_345_67_89,
        "oktatas_datuma": "lsd.csatolt lista.",# lásd csatolt lista vagy "ha egy jog van akkor is lesz csatolt lista 1 jog kérvény lesz benne. MINDENKÉPPEN LISTÁZD!!!"
        "oktatasi_azonosito": "lsd.csatolt lista",# lásd csatolt lista vagy  Üresen marad!!!!
        "rendszer_adminisztrator":"",#Üresen maradó.
        }


#Létrehozom a dokumentumot:
document = Document()

table = document.add_table(rows=15,cols=3,style='TableGrid')
#Beállítom a betű típust:
style = document.styles["Normal"]
style.font.name = "Calibri"

#Az első hatot kell  mergelem:
for row in range(0,6):
    cell = table.cell(row,0)
    for col in range(1,3):
        cell.merge(table.cell(row,col))

#Fejléc megadása:
table.cell(0,0).text = "Felhasználói jogok igénylése"


# Adat feltöltés első hatnál:
for row_count in range(1,6):
    table.cell(row_count,0).text = list(test_data.keys())[int(row_count)-1]

#2 oszlopos cellák:
cell = table.cell(6,0)
cell.merge(table.cell(6,1))

cell = table.cell(7,0)
cell.merge(table.cell(7,1))
       


  



#Fájl mentése:
document.save(str("Kérvény")+ '_' + str(test_data["torzs_szama"])+ '_' + test_data["felhasznalo"]+ '.docx')
#Funckió: